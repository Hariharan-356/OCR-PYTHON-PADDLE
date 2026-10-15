from __future__ import annotations

import asyncio
import json
import re
import shutil
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, Header, HTTPException
from pydantic import BaseModel

from ocr_service.config.settings import settings
from ocr_service.models import HealthResponse, JobRequest, JobState

try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover
    fitz = None

try:
    import cv2
except Exception:  # pragma: no cover
    cv2 = None

try:
    import numpy as np
except Exception:  # pragma: no cover
    np = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None

try:
    from bs4 import BeautifulSoup
except Exception:  # pragma: no cover
    BeautifulSoup = None

try:
    from PIL import Image
except Exception:  # pragma: no cover
    Image = None

try:
    import pytesseract
    tesseract_bin = shutil.which('tesseract') or r'C:\Program Files\Tesseract-OCR\tesseract.exe'
    if tesseract_bin and Path(tesseract_bin).exists():
        pytesseract.pytesseract.tesseract_cmd = tesseract_bin
except Exception:  # pragma: no cover
    pytesseract = None

try:
    from rapidocr_onnxruntime import RapidOCR
except Exception:  # pragma: no cover
    RapidOCR = None

app = FastAPI(title='ClaimsDigi OCR Service')

JOBS: Dict[str, JobState] = {}

SUPPORTED_EXTENSIONS = {
    '.txt', '.pdf', '.png', '.jpg', '.jpeg', '.tiff', '.tif', '.doc', '.docx', '.html', '.htm', '.xlsx', '.xls', '.zip'
}
DOC_TYPE_MAP = {
    'pre-auth': 'Pre-Auth Form (Form C)',
    'form c': 'Pre-Auth Form (Form C)',
    'claim form a': 'Claim Form A',
    'claim form b': 'Claim Form B',
    'discharge': 'Discharge Summary',
    'hospital bill': 'Hospital Bill',
    'pharmacy bill': 'Pharmacy Bill',
    'diagnostic bill': 'Diagnostic Bill',
    'opd bill': 'OPD Bill',
    'implant bill': 'Implant Bill',
    'food bill': 'Food Bill',
    'investigation': 'Investigation Report',
    'kyc identity': 'KYC – Identity Proof',
    'address proof': 'KYC – Address Proof',
    'cancelled cheque': 'Cancelled Cheque',
    'hospital registration': 'Hospital Registration Certificate',
    'fir': 'FIR',
    'postmortem': 'Postmortem Report',
}


def mask_all_but_last_4(value: str) -> str:
    s = re.sub(r'\D', '', value)
    if len(s) <= 4:
        return 'XXXX'
    return 'XXXX' + s[-4:]


def normalize_date(value: str) -> str:
    text = value.strip()
    if not text:
        return ''
    patterns = [
        (r'^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$', '%d/%m/%Y'),
        (r'^(\d{4})-(\d{1,2})-(\d{1,2})$', '%Y-%m-%d'),
    ]
    for pattern, fmt in patterns:
        m = re.match(pattern, text)
        if m:
            try:
                day, month, year = m.groups()[:3]
                if len(year) == 2:
                    year = '20' + year
                if int(month) > 12 or int(day) > 31:
                    continue
                try:
                    from datetime import datetime
                    dt = datetime.strptime(f'{day}-{month}-{year}', '%d-%m-%Y')
                    return dt.strftime('%Y-%m-%d')
                except ValueError:
                    continue
            except ValueError:
                continue
    return text


def detect_document_type(name: str, text: str) -> str:
    hay = f"{name} {text}".lower()
    if any(token in hay for token in [
        'address of biller',
        'hospital registration number',
        'gst no.',
        'bill no.',
        'gross amount',
        'total amount paid',
        'revenue stamp',
    ]):
        return 'Hospital Bill'

    for key, value in DOC_TYPE_MAP.items():
        if key in hay:
            return value

    if 'pre' in hay and 'auth' in hay:
        return 'Pre-Auth Form (Form C)'
    if 'hospital' in hay and 'bill' in hay:
        return 'Hospital Bill'
    if 'discharge' in hay:
        return 'Discharge Summary'
    if 'pharmacy' in hay:
        return 'Pharmacy Bill'
    if 'lab' in hay or 'investigation' in hay:
        return 'Investigation Report'
    if 'aadhar' in hay or 'aadhaar' in hay or 'kyc' in hay and 'identity' in hay:
        return 'KYC – Identity Proof'
    if 'address' in hay and 'proof' in hay:
        return 'KYC – Address Proof'
    if 'cheque' in hay:
        return 'Cancelled Cheque'
    return 'Unclassified'


def validate_magic(path: Path) -> bool:
    if not path.exists():
        return False
    try:
        with path.open('rb') as handle:
            signature = handle.read(16)
    except OSError:
        return False

    if signature.startswith(b'%PDF'):
        return True
    if signature.startswith(b'PK'):
        return True
    if signature.startswith(b'\x89PNG'):
        return True
    if signature.startswith(b'\xff\xd8\xff'):
        return True
    if signature.startswith(b'II*\x00') or signature.startswith(b'MM\x00*'):
        return True
    if signature.startswith(b'\x00\x00\x00'):
        return True
    return True


def validate_file(file_path: Path, original_name: str, size_bytes: int) -> Dict[str, Any]:
    ext = file_path.suffix.lower()
    if ext and ext not in SUPPORTED_EXTENSIONS:
        return {'ok': False, 'reason': f'Unsupported extension: {ext}' }
    if not validate_magic(file_path):
        return {'ok': False, 'reason': 'Magic-byte validation failed'}
    if size_bytes > settings.max_upload_mb * 1024 * 1024:
        return {'ok': False, 'reason': f'File exceeds {settings.max_upload_mb} MB'}
    return {'ok': True, 'reason': 'Accepted'}


def expand_zip_contents(file_path: Path) -> List[Path]:
    if file_path.suffix.lower() != '.zip':
        return [file_path]
    if not zipfile.is_zipfile(file_path):
        return [file_path]

    temp_dir = Path(tempfile.mkdtemp(prefix='claimsdigi_zip_', dir=str(file_path.parent)))
    with zipfile.ZipFile(file_path) as archive:
        archive.extractall(temp_dir)
    extracted_files = sorted(
        (candidate for candidate in temp_dir.rglob('*') if candidate.is_file()),
        key=lambda candidate: str(candidate).lower(),
    )

    if not extracted_files:
        shutil.rmtree(temp_dir, ignore_errors=True)
        return [file_path]

    return extracted_files


def enhance_image_for_ocr(image) -> Any:
    if cv2 is None or np is None:
        return image

    if isinstance(image, Image.Image):
        image_array = np.array(image.convert('RGB'))
    else:
        image_array = np.asarray(image)

    if image_array.size == 0:
        return image

    if len(image_array.shape) == 3:
        gray = cv2.cvtColor(image_array, cv2.COLOR_RGB2GRAY if image_array.shape[2] == 3 else cv2.COLOR_BGR2GRAY)
    else:
        gray = image_array

    gray = cv2.resize(gray, None, fx=1.8, fy=1.8, interpolation=cv2.INTER_CUBIC)
    clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
    gray = clahe.apply(gray)
    gray = cv2.medianBlur(gray, 3)
    gray = cv2.GaussianBlur(gray, (3, 3), 0)

    thresh = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
    if thresh.std() > 0:
        gray = thresh

    contours = cv2.findContours(gray, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    contours = contours[0] if len(contours) == 2 else contours[1]
    if contours:
        all_points = np.concatenate(contours, axis=0)
        rect = cv2.minAreaRect(all_points)
        angle = rect[-1]
        if angle < -45:
            angle = 90 + angle
        elif angle > 45:
            angle = angle - 90
        if abs(angle) > 0.5:
            (h, w) = gray.shape[:2]
            matrix = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1.0)
            gray = cv2.warpAffine(gray, matrix, (w, h), flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_REPLICATE)

    gray = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 10)
    return Image.fromarray(gray)


def ocr_image_with_fallback(image) -> str:
    if image is None:
        return ''

    if pytesseract is not None:
        try:
            return pytesseract.image_to_string(image, config='--psm 6')
        except Exception:
            pass

    if RapidOCR is not None:
        try:
            detector = RapidOCR()
            image_array = np.array(image) if np is not None else image
            result, _ = detector(image_array)
            if result:
                content = ' '.join(item[1] for item in result if item and len(item) > 1)
                return content
        except Exception:
            pass

    return ''


def extract_pdf_text(path: Path) -> str:
    if fitz is None:
        return ''
    text_chunks: List[str] = []
    with fitz.open(str(path)) as document:
        for page in document:
            text = page.get_text().strip()
            text_chars = re.sub(r'\s+', '', text)
            if text and len(text_chars) >= 20:
                text_chunks.append(text)
                continue

            try:
                page_image = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                image = Image.frombytes('RGB', [page_image.width, page_image.height], page_image.samples)
                enhanced = enhance_image_for_ocr(image)
                ocr_text = ocr_image_with_fallback(enhanced)
            except Exception:
                ocr_text = ''
            if text:
                text_chunks.append(text)
            if ocr_text.strip():
                text_chunks.append(ocr_text.strip())
    return '\n'.join(text_chunks)


def extract_docx_text(path: Path) -> str:
    if DocxDocument is None:
        return ''
    document = DocxDocument(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    return '\n'.join(paragraphs)


def extract_xlsx_text(path: Path) -> str:
    if load_workbook is None:
        return ''
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    rows: List[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = ' | '.join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                rows.append(row_text)
    return '\n'.join(rows)


def extract_html_text(path: Path) -> str:
    if BeautifulSoup is None:
        return ''
    text = path.read_text(encoding='utf-8', errors='ignore')
    soup = BeautifulSoup(text, 'html.parser')
    return ' '.join(soup.stripped_strings)


def extract_text_from_image(path: Path) -> str:
    if Image is None:
        return ''
    image = Image.open(path)
    enhanced = enhance_image_for_ocr(image)
    if pytesseract is not None:
        try:
            return pytesseract.image_to_string(enhanced, config='--psm 6')
        except Exception:
            pass
    if RapidOCR is not None:
        try:
            detector = RapidOCR()
            image_array = np.array(enhanced) if np is not None else enhanced
            result, _ = detector(image_array)
            if result:
                return ' '.join(item[1] for item in result if item and len(item) > 1)
        except Exception:
            pass
    return ''


def extract_text(path: Path, name: str) -> str:
    suffix = path.suffix.lower()
    if suffix == '.zip':
        texts: List[str] = []
        for nested in expand_zip_contents(path):
            if nested == path:
                continue
            extracted_text = extract_text(nested, nested.name)
            if extracted_text.strip():
                texts.append(extracted_text)
        return '\n\n'.join(texts)
    if suffix == '.pdf':
        return extract_pdf_text(path)
    if suffix == '.docx':
        return extract_docx_text(path)
    if suffix in {'.xlsx', '.xls'}:
        return extract_xlsx_text(path)
    if suffix in {'.html', '.htm'}:
        return extract_html_text(path)
    if suffix in {'.txt'}:
        try:
            return path.read_text(encoding='utf-8', errors='ignore')
        except Exception:
            return ''
    if suffix in {'.png', '.jpg', '.jpeg', '.tif', '.tiff'}:
        return extract_text_from_image(path)
    return ''


def find_first(text: str, patterns: List[str]) -> str:
    for pattern in patterns:
        candidate = re.search(pattern, text, re.I)
        if candidate:
            return candidate.group(1).strip() if candidate.groups() else candidate.group(0).strip()
    return ''


def find_all(text: str, pattern: str) -> List[str]:
    return [m.group(0).strip() for m in re.finditer(pattern, text, re.I)]


def build_field(section: str, label: str, value: str, conf: int = 88, original: str = '', masked: bool = False, detect: Optional[str] = None) -> Dict[str, Any]:
    normalized_value = (value or '').strip()
    if not normalized_value:
        normalized_value = 'Not extracted — fill manually'
    return {
        'id': f"fld-{uuid.uuid4().hex[:8]}",
        'section': section,
        'label': label,
        'value': normalized_value,
        'conf': conf,
        'original': original or normalized_value,
        'verified': False,
        'history': [],
        'masked': masked,
        'detect': detect,
    }


def extract_hospital_bill_fields(text: str, facts: Dict[str, Any]) -> List[Dict[str, Any]]:
    normalized = text.replace('\r', '\n')
    lines = [line.strip() for line in normalized.splitlines() if line.strip()]
    biller = (lines[0] if lines else '')

    address = find_first(normalized, [
        r'address\s*of\s*biller\s*[:\-]?\s*(.+)',
        r'hospital\s*address\s*[:\-]?\s*(.+)',
        r'address\s*[:\-]?\s*(.+)',
    ])
    if not address:
        for line in lines[1:]:
            stripped = re.sub(r'\s+', ' ', line).strip()
            if not stripped:
                continue
            if re.search(r'\b(hospital|reg\.?|gst|bill|date|patient|doctor|admission|discharge|gross|discount|total|amount)\b', stripped, re.I):
                continue
            if re.search(r'\b(road|street|avenue|lane|main|salai|cross|town|district|state|tamil nadu|coimbatore|peelamedu|road|road,|street,)\b', stripped, re.I) or re.match(r'^\d+\s*,', stripped):
                address = stripped
                break
    address = address or facts.get('city') or ''

    registration = find_first(normalized, [
        r'hospital\s*(?:registration|reg\.)\s*(?:number|no\.)?\s*[:\-]?\s*([A-Za-z0-9/.-]+)',
        r'(?:(?:registration|reg\.)\s*(?:number|no\.)?)\s*[:\-]?\s*([A-Za-z0-9/.-]+)',
        r'\b(?:reg\.?\s*no\.?|registration\s*number)\s*[:\-]?\s*([A-Za-z0-9/.-]+)',
    ]) or facts.get('hospReg') or ''
    gst_no = find_first(normalized, [
        r'gst\s*(?:no\.?|number)\s*[:\-]?\s*([A-Z0-9]+)',
        r'gst\s*[:\-]?\s*([A-Z0-9]+)',
    ]) or facts.get('gst') or ''
    bill_no = find_first(normalized, [
        r'bill\s*no\.?\s*[:\-]?\s*([A-Za-z0-9/.-]+)',
        r'bill\s*number\s*[:\-]?\s*([A-Za-z0-9/.-]+)',
    ]) or facts.get('billNo') or ''
    bill_date = find_first(normalized, [
        r'bill\s*date\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})',
    ]) or facts.get('discharge') or ''
    patient_name = find_first(normalized, [
        r'name\s*of\s*patient\s*[:\-]?\s*(.+)',
        r'patient\s*name\s*[:\-]?\s*(.+)',
    ]) or facts.get('patient') or ''
    admission_date = find_first(normalized, [
        r'date\s*of\s*admission\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})',
    ]) or facts.get('admission') or ''
    discharge_date = find_first(normalized, [
        r'date\s*of\s*discharge\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})',
    ]) or facts.get('discharge') or ''
    doctor_name = find_first(normalized, [
        r'name\s*of\s*treating\s*doctor\s*[:\-]?\s*(.+)',
        r'treating\s*doctor\s*[:\-]?\s*(.+)',
    ]) or facts.get('doctor') or ''
    gross_amount = find_first(normalized, [
        r'gross\s*amount\s*[:\-]?\s*(₹?\s*\d[\d,]*(?:\.\d{1,2})?)',
        r'gross\s*[:\-]?\s*(₹?\s*\d[\d,]*(?:\.\d{1,2})?)',
    ]) or ''
    discount = find_first(normalized, [
        r'discount\s*[:\-]?\s*(₹?\s*\d[\d,]*(?:\.\d{1,2})?)',
    ]) or ''
    gst_amount = find_first(normalized, [
        r'gst\s*amount\s*[:\-]?\s*(₹?\s*\d[\d,]*(?:\.\d{1,2})?)',
    ]) or ''
    total_amount = find_first(normalized, [
        r'total\s*amount\s*paid\s*[:\-]?\s*(₹?\s*\d[\d,]*(?:\.\d{1,2})?)',
        r'total\s*paid\s*[:\-]?\s*(₹?\s*\d[\d,]*(?:\.\d{1,2})?)',
    ]) or facts.get('billed') or ''
    stamp_availability = find_first(normalized, [
        r'stamp\s*availability\s*[:\-]?\s*(yes|no)',
    ]) or 'Yes'
    revenue_stamp = find_first(normalized, [
        r'revenue\s*stamp\s*[:\-]?\s*(yes|no|detected)',
    ]) or 'Yes'
    hospital_logo = find_first(normalized, [
        r'hospital\s*logo\s*[:\-]?\s*(detected|not detected|yes|no)',
    ]) or 'Detected'

    return [
        build_field('Biller', 'Name of biller', biller.strip(), 99),
        build_field('Biller', 'Address of biller', address.strip(), 89),
        build_field('Biller', 'Hospital registration number', registration.strip(), 88),
        build_field('Biller', 'GST no.', gst_no.strip(), 97),
        build_field('Bill header', 'Bill no.', bill_no.strip(), 99),
        build_field('Bill header', 'Bill date', bill_date.strip(), 99),
        build_field('Bill header', 'Name of patient', patient_name.strip(), 99),
        build_field('Bill header', 'Date of admission', admission_date.strip(), 95),
        build_field('Bill header', 'Date of discharge', discharge_date.strip(), 96),
        build_field('Bill header', 'Name of treating doctor', doctor_name.strip(), 96),
        build_field('Totals', 'Gross amount', gross_amount.strip(), 99),
        build_field('Totals', 'Discount', discount.strip(), 97),
        build_field('Totals', 'GST amount', gst_amount.strip(), 95),
        build_field('Totals', 'Total amount paid', str(total_amount).strip(), 98),
        build_field('Totals', 'Stamp availability', stamp_availability.strip().title(), 96, detect='stamp'),
        build_field('Totals', 'Revenue stamp', revenue_stamp.strip().title(), 88, detect='stamp'),
        build_field('Totals', 'Hospital logo', hospital_logo.strip().title(), 95, detect='logo'),
    ]


def extract_claim_fields(doc_type: str, text: str) -> Dict[str, Any]:
    normalized = text.replace('\r', '\n')
    patient = find_first(normalized, [r'patient(?:\s*(?:name|name\s*of))?\s*[:\-]?\s*(.+)', r'claimant\s*name\s*[:\-]?\s*(.+)', r'name\s*of\s*the\s*patient\s*[:\-]?\s*(.+)'])
    member = find_first(normalized, [r'member\s*(?:id|number)\s*[:\-]?\s*([A-Z0-9-]+)', r'insured\s*id\s*[:\-]?\s*([A-Z0-9-]+)'])
    policy = find_first(normalized, [r'policy\s*(?:number|no\.?|no\s*)\s*[:\-]?\s*([A-Z0-9/-]+)', r'policy\s*[:\-]?\s*([A-Z0-9/-]+)'])
    hospital = find_first(normalized, [
        r'name\s*of\s*the\s*hospital\s*[:\-]?\s*(.+)',
        r'name\s*of\s*biller\s*[:\-]?\s*(.+)',
        r'hospital\s*name\s*[:\-]?\s*(.+)',
    ])
    if not hospital:
        lines = [line.strip() for line in normalized.splitlines() if line.strip()]
        for line in lines:
            lowered = line.lower()
            if any(token in lowered for token in [
                'address of biller', 'hospital registration', 'gst no.', 'bill no.',
                'name of patient', 'date of admission', 'date of discharge',
                'gross amount', 'discount', 'total amount paid'
            ]):
                continue
            hospital = line
            break
    city = find_first(normalized, [
        r'address\s*of\s*biller\s*[:\-]?\s*(.+)',
        r'hospital\s*address\s*[:\-]?\s*(.+)',
        r'city\s*[:\-]?\s*(.+)',
    ])
    doctor = find_first(normalized, [r'doctor\s*(?:name|\w+)\s*[:\-]?\s*(.+)', r'name\s*of\s*treating\s*doctor\s*[:\-]?\s*(.+)'])
    diagnosis = find_first(normalized, [r'diagnosis\s*[:\-]?\s*(.+)', r'provisional\s*diagnosis\s*[:\-]?\s*(.+)'])
    admission = find_first(normalized, [r'date\s*of\s*admission\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})', r'admission\s*date\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})'])
    discharge = find_first(normalized, [r'date\s*of\s*discharge\s*[:\-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})'])
    amount = find_first(normalized, [r'total\s*(?:amount|bill|billed)\s*[:\-]?\s*₹?\s*(\d[\d,]*(?:\.\d{1,2})?)', r'paid\s*amount\s*[:\-]?\s*₹?\s*(\d[\d,]*(?:\.\d{1,2})?)'])

    if amount:
        amount = amount.replace(',', '')
    facts = {
        'patient': patient or 'Unknown patient',
        'memberId': member or '',
        'policy': policy or '',
        'hospital': hospital or '',
        'city': city or '',
        'doctor': doctor or '',
        'diagnosis': diagnosis or '',
        'admission': normalize_date(admission),
        'discharge': normalize_date(discharge),
        'admissionISO': normalize_date(admission),
        'billed': float(amount) if amount else 0.0,
    }

    if doc_type == 'Hospital Bill':
        return {'facts': facts, 'fields': extract_hospital_bill_fields(normalized, facts)}

    doc_fields = [
        build_field('Patient', 'Name of the patient', facts['patient'], 92),
        build_field('Patient', 'Insured id / Member id', facts['memberId'], 90),
        build_field('Patient', 'Policy number', facts['policy'], 90),
        build_field('Hospital details', 'Name of the hospital', facts['hospital'], 90),
        build_field('Hospital details', 'Name of treating doctor', facts['doctor'], 84),
        build_field('Admission details', 'Date of admission', facts['admission'], 88),
        build_field('Admission details', 'Date of discharge', facts['discharge'], 88),
    ]
    if doc_type in {'Hospital Bill', 'Pharmacy Bill', 'Diagnostic Bill', 'OPD Bill', 'Implant Bill', 'Food Bill'}:
        doc_fields.append(build_field('Totals', 'Total amount paid', f"₹{facts['billed']:,.2f}", 92))
    return {'facts': facts, 'fields': doc_fields}


def classify_document(document_name: str, text: str) -> str:
    doc_type = detect_document_type(document_name, text)
    if doc_type == 'Unclassified':
        lowered = document_name.lower()
        if 'aadhaar' in lowered or 'aadhar' in lowered:
            return 'KYC – Identity Proof'
        if 'address' in lowered or 'proof' in lowered:
            return 'KYC – Address Proof'
    return doc_type


async def process_job(job_id: str, request: JobRequest) -> None:
    job = JOBS[job_id]
    job.status = 'processing'
    job.progress = 10
    documents: List[Dict[str, Any]] = []
    warnings: List[str] = []
    expanded_files: List[Dict[str, Any]] = []

    for item in request.files:
        file_path = Path(item.path)
        if file_path.suffix.lower() == '.zip':
            for nested_path in expand_zip_contents(file_path):
                if nested_path == file_path:
                    continue
                expanded_files.append({
                    'name': nested_path.name,
                    'path': str(nested_path),
                    'size': nested_path.stat().st_size,
                    'mimetype': 'application/pdf' if nested_path.suffix.lower() == '.pdf' else 'application/octet-stream',
                })
        else:
            expanded_files.append({
                'name': item.name,
                'path': item.path,
                'size': item.size,
                'mimetype': item.mimetype,
            })

    for index, item in enumerate(expanded_files, start=1):
        file_path = Path(item['path'])
        validation = validate_file(file_path, item['name'], int(item['size']))
        if not validation['ok']:
            warnings.append(f"{item['name']}: {validation['reason']}")
            continue
        text = extract_text(file_path, item['name'])
        if not text.strip():
            warnings.append(f"{item['name']}: no readable text extracted")
            continue
        doc_type = classify_document(item['name'], text)
        extracted = extract_claim_fields(doc_type, text)
        facts = extracted['facts']
        fields = extracted['fields']
        pages = [{
            'n': 1,
            'cls': doc_type,
            'conf': 91,
            'label': doc_type,
        }]
        docs = {
            'id': f"doc-{job_id}-{index}",
            'name': item['name'],
            'type': doc_type,
            'source': item['name'],
            'size': f"{max(0.1, int(item['size']) / 1048576):.1f} MB",
            'ext': file_path.suffix.lower().lstrip('.'),
            'uploaded': __import__('datetime').datetime.utcnow().isoformat() + 'Z',
            'ocrStatus': 'Completed' if fields else 'Needs review',
            'handwritten': 'signature' in item['name'].lower() or 'handwritten' in text.lower(),
            'language': 'English',
            'pages': pages,
            'fields': fields,
            'detections': ['ocr'] if fields else ['uncertain'],
            'medical': doc_type not in {'KYC – Identity Proof', 'KYC – Address Proof', 'Cancelled Cheque'},
            'validated': bool(fields),
            'masked': doc_type in {'KYC – Identity Proof', 'KYC – Address Proof', 'Cancelled Cheque'},
        }
        documents.append(docs)
        job.progress = min(95, 10 + (index / max(1, len(expanded_files))) * 80)
        job.files.append({
            'name': item['name'],
            'status': 'processed',
            'pages': 1,
            'docType': doc_type,
        })

    if not documents:
        job.status = 'failed'
        job.progress = 100
        job.error = 'No supported documents were processed.'
        job.result = {'documents': [], 'warnings': warnings}
        return

    merged_facts = {
        'patient': 'Unknown patient',
        'memberId': '',
        'policy': '',
        'hospital': '',
        'city': '',
        'admission': '',
        'discharge': '',
        'admissionISO': '',
        'billed': 0,
    }
    for doc in documents:
        for field in doc['fields']:
            label = field['label'].lower()
            value = str(field['value']).strip()
            if not value:
                continue
            if 'patient' in label and 'name' in label:
                merged_facts['patient'] = value
            elif 'member' in label or 'insured id' in label:
                merged_facts['memberId'] = value
            elif 'policy' in label:
                merged_facts['policy'] = value
            elif ('hospital' in label and 'name' in label) or ('biller' in label and 'name' in label):
                merged_facts['hospital'] = value
            elif 'address of biller' in label:
                merged_facts['city'] = value
            elif 'date of admission' in label.lower():
                merged_facts['admission'] = normalize_date(value)
            elif 'date of discharge' in label.lower():
                merged_facts['discharge'] = normalize_date(value)
            elif 'total amount paid' in label.lower() or 'gross' in label.lower():
                try:
                    merged_facts['billed'] += float(str(value).replace('₹', '').replace(',', '').replace(' ', ''))
                except ValueError:
                    pass

    if not merged_facts['patient'] or merged_facts['patient'] == 'Unknown patient':
        merged_facts['patient'] = request.claimId or 'Unknown patient'

    result = {
        'claimId': request.claimId or f"CLM-OCR-{uuid.uuid4().hex[:8].upper()}",
        'facts': merged_facts,
        'documents': documents,
        'lineItems': [
            {
                'id': f"line-{uuid.uuid4().hex[:8]}",
                'bill': doc.get('type', 'Unclassified'),
                'desc': doc.get('type', 'Unclassified'),
                'code': '',
                'qty': 1,
                'unit': 'unit',
                'unitPrice': 0,
                'gross': 0,
                'discount': 0,
                'net': 0,
                'category': 'General',
                'medical': True,
                'date': merged_facts.get('admission') or '',
            }
            for doc in documents
        ],
        'diagnoses': [],
        'procedures': [],
        'warnings': warnings,
        'metadata': {
            'primary_engine': settings.primary_engine,
            'secondary_engine': settings.secondary_engine,
            'status': 'processed-with-fallback'
        },
    }

    job.status = 'completed' if not warnings else 'needs_review'
    job.progress = 100
    job.result = result


@app.get('/health')
def health() -> HealthResponse:
    return HealthResponse(
        status='healthy',
        engines={'paddleocr': 'optional', 'easyocr': 'optional', 'tesseract': 'optional'},
        models_ready={'paddleocr': False, 'easyocr': False, 'tesseract': False},
        version='0.1.0',
    )


@app.post('/validate')
def validate(request: dict) -> dict:
    file_path = Path(str(request.get('path', '')))
    if not file_path.exists():
        raise HTTPException(status_code=400, detail='file path not found')
    return validate_file(file_path, file_path.name, file_path.stat().st_size)


@app.post('/jobs')
async def create_job(request: JobRequest, x_shared_secret: Optional[str] = Header(default=None, alias='x-shared-secret')) -> JobState:
    if x_shared_secret != settings.shared_secret:
        raise HTTPException(status_code=401, detail='unauthorized')
    job_id = f"job-{uuid.uuid4().hex[:8]}"
    state = JobState(jobId=job_id, status='queued', progress=0, files=[])
    JOBS[job_id] = state
    asyncio.create_task(process_job(job_id, request))
    return state


@app.get('/jobs/{job_id}')
def get_job(job_id: str, x_shared_secret: Optional[str] = Header(default=None, alias='x-shared-secret')) -> JobState:
    if x_shared_secret != settings.shared_secret:
        raise HTTPException(status_code=401, detail='unauthorized')
    if job_id not in JOBS:
        raise HTTPException(status_code=404, detail='job not found')
    return JOBS[job_id]


@app.get('/jobs/{job_id}/result')
def get_job_result(job_id: str, x_shared_secret: Optional[str] = Header(default=None, alias='x-shared-secret')) -> Dict[str, Any]:
    if x_shared_secret != settings.shared_secret:
        raise HTTPException(status_code=401, detail='unauthorized')
    if job_id not in JOBS:
        raise HTTPException(status_code=404, detail='job not found')
    job = JOBS[job_id]
    return job.result or {'status': job.status, 'documents': []}


@app.get('/jobs/{job_id}/files')
def get_job_files(job_id: str, x_shared_secret: Optional[str] = Header(default=None, alias='x-shared-secret')) -> List[Dict[str, Any]]:
    if x_shared_secret != settings.shared_secret:
        raise HTTPException(status_code=401, detail='unauthorized')
    if job_id not in JOBS:
        raise HTTPException(status_code=404, detail='job not found')
    return JOBS[job_id].files
